#!/usr/bin/env python3
"""
importar_estructuras_excel.py
==============================
Importador directo del Excel de Estructuras de Costos de YEIKAR
(CAMAS CON SUS MESAS DE NOCHE.xlsx).

Diferencia clave con la versión previa:
  - Clasifica las filas de cada hoja por SECCIÓN (EBANISTERIA, TENDIDO, COLA_DE_PATO, TAPICERIA, PINTURA, TERMINACION, NOCHEROS, MANO_DE_OBRA).
  - Descarta explícitamente filas calculadas (SUBTOTALES, TOTALES, RECARGOS % GASTOS DE...).
  - Identifica filas de MANO DE OBRA y las asigna a la sección MANO_DE_OBRA.
  - Permite repetición legítima de materiales entre distintas secciones (fix Bug 2: filtra por producto_id + material_id + seccion).
  - Soporta opción --dry-run para validar la clasificación antes de escribir en la base de datos.
"""

import sys
import re
import argparse
import traceback
from pathlib import Path
from decimal import Decimal, InvalidOperation

sys.path.append(str(Path(__file__).parent.parent))

import openpyxl
from app.db.session import session_local
from app.modules.productos.model import Producto, ProductoMaterial, Material, ReglaGastoSeccion
from app.modules.catalogos.model import TipoProducto, UnidadMedida

EXCEL_PATH = (
    Path(__file__).parent.parent.parent
    / "docs/info_para_usar/estructuras/CAMAS CON SUS MESAS DE NOCHE.xlsx"
)

# Diccionario de encabezados de sección en las hojas de Excel
HEADERS_SECCION = {
    "SECCION TAPICERIA": "TAPICERIA",
    "SECCION TAPICERÍA": "TAPICERIA",
    "SECCION TERMINACION": "TERMINACION",
    "SECCION TERMINACIÓN": "TERMINACION",
    "SECCION PINTURA": "PINTURA",
    "TENDIDO DE REJAS": "TENDIDO",
    "TENDIDOS": "TENDIDO",
    "TENDIDO": "TENDIDO",
    "COLA DE PATO": "COLA_DE_PATO",
    "NOCHEROS": "NOCHEROS",
    "NOCHERO": "NOCHEROS",
    "SECCION EBANISTERIA": "EBANISTERIA",
    "SECCION EBANISTERÍA": "EBANISTERIA",
}

# Prefijos o palabras clave que indican fila CALCULADA (subtotal/total/gasto/precio) -> NUNCA insertar como material
PREFIJOS_IGNORAR = (
    "TOTAL", "SUBTOTAL", "SUB TOTAL", "COSTO DE", "GASTOS DE", "GASTOS", "gastos de",
    "TOTAL COSTO DE PRODUCCION", "TOTAL COSTO DE PRODUCCIÓN", "TOTAL PRODUCCION", "TOTAL PRODUCCIÓN",
    "PRECIO DE VENTA", "Precio de Venta", "IVA", "TOTAL A PAGAR", "Total a Pagar",
    "MÁS GANANCIA", "MAS GANANCIA", "Más Ganancia", "MATERIA PRIMA",
    "COMERCIALIZADORA", "ESTRUCTURA", "RIF", "NOMBRE DEL PRODUCTO", "FECHA",
    "EBANISTA", "FABRICO", "FABRICÓ", "VALOR DE", "CONCEPTO"
)

# Palabras clave que indican MANO DE OBRA, no insumo físico
KEYWORDS_MANO_OBRA = (
    "Tabulación", "TABULACIÓN", "HECHURA", "m.o", "M.O", "PAGO AL CONTRATISTA",
    "Pago al Contratista", "PAGO AL EBANISTA", "ebanisteria 300", "ebanisteria", "ebanisteria X",
    "POSTURA DE", "PAGO DE"
)

UNIDADES_NORMALIZADAS = {
    "CMS": ("cm", "Centímetros"),
    "METROS": ("m", "Metros"),
    "METRO": ("m", "Metros"),
    "LITRO": ("L", "Litros"),
    "LITROS": ("L", "Litros"),
    "LAMINA": ("lam", "Láminas"),
    "LAMINAS": ("lam", "Láminas"),
    "CAJAS": ("cj", "Cajas"),
    "CAJA": ("cj", "Cajas"),
    "UNIDAD": ("und", "Unidades"),
    "UNIDADES": ("und", "Unidades"),
    "PAR": ("par", "Pares"),
    "TIRA": ("tira", "Tiras"),
    "TIRAS": ("tira", "Tiras"),
    "PEOPLE": ("people", "People"),
    "MANO DE OBRA": ("mo", "Mano de Obra"),
}


def limpiar_nombre_hoja(nombre: str) -> str:
    if not nombre:
        return "SIN NOMBRE"
    nombre = re.sub(r'\b\d{1,2}[-/]\d{1,2}[-/]\d{2,4}\b', '', nombre)
    nombre = re.sub(r'\d{1,2}//\d{2}/\d{4}', '', nombre)
    nombre = re.sub(r'\s*\(\d+\)\s*', '', nombre)
    nombre = ' '.join(nombre.split()).strip().strip('.,- ')
    return nombre[:150] if nombre else "SIN NOMBRE"


def clasificar_fila(nombre_fila: str) -> str:
    """
    Devuelve tipo de fila:
      'header_seccion' | 'ignorar' | 'mano_obra' | 'material'
    """
    texto = nombre_fila.strip()
    texto_upper = texto.upper()

    # 1. ¿Es un encabezado explícito de sección?
    for header in HEADERS_SECCION:
        if header.lower() in texto.lower():
            return "header_seccion"

    # 2. ¿Es una fila calculada / subtotal / gasto / total?
    for prefijo in PREFIJOS_IGNORAR:
        if texto_upper.startswith(prefijo) or texto_upper == prefijo:
            return "ignorar"

    # 3. ¿Es mano de obra?
    for kw in KEYWORDS_MANO_OBRA:
        if kw.lower() in texto.lower():
            return "mano_obra"

    return "material"


def parsear_hoja(ws, nombre_producto: str) -> dict:
    receta_filas = []
    seccion_actual = "EBANISTERIA"  # Default al inicio de cada hoja
    precio_costo = None
    precio_venta = None
    precio_con_iva = None

    for row in ws.iter_rows(values_only=True):
        celdas = list(row)
        if not any(v is not None for v in celdas):
            continue

        primera_col = celdas[0]
        fila_txt = " ".join(str(v) for v in celdas if v is not None).upper()
        nums = [v for v in celdas if isinstance(v, (int, float)) and v > 0]

        # Extraer precios globales si aparecen en la hoja
        if "TOTAL PRODUCCION" in fila_txt or "TOTAL PRODUCCIÓN" in fila_txt:
            if nums and precio_costo is None:
                precio_costo = Decimal(str(nums[0]))
            continue

        if "PRECIO DE VENTA DEL PRODUCTO" in fila_txt or ("PRECIO DE VENTA" in fila_txt and "PRODUCTO" in fila_txt):
            if nums and precio_venta is None:
                precio_venta = Decimal(str(nums[0]))
            continue

        if "TOTAL A PAGAR" in fila_txt:
            if nums and precio_con_iva is None:
                precio_con_iva = Decimal(str(nums[0]))
            continue

        if not primera_col or not isinstance(primera_col, str):
            continue

        col_a_str = primera_col.strip()
        if len(col_a_str) < 2:
            continue

        tipo = clasificar_fila(col_a_str)

        if tipo == "header_seccion":
            for header, sec_target in HEADERS_SECCION.items():
                if header.lower() in col_a_str.lower():
                    seccion_actual = sec_target
                    break
            continue

        if tipo == "ignorar":
            continue

        # Obtener precio total de la fila si existe
        precio_total = Decimal(0)
        if len(celdas) > 5 and isinstance(celdas[5], (int, float)) and celdas[5] > 0:
            precio_total = Decimal(str(celdas[5]))
        elif len(celdas) > 4 and isinstance(celdas[4], (int, float)) and celdas[4] > 0:
            if len(celdas) == 5:
                precio_total = Decimal(str(celdas[4]))

        # Obtener cantidad y v_unit
        qty = celdas[1] if len(celdas) > 1 else None
        v_unit_val = None
        if len(celdas) > 4 and isinstance(celdas[4], (int, float)) and celdas[4] > 0:
            v_unit_val = celdas[4]
        elif len(celdas) > 3 and isinstance(celdas[3], (int, float)) and celdas[3] > 0:
            v_unit_val = celdas[3]

        try:
            cantidad_dec = Decimal(str(qty)) if qty is not None and isinstance(qty, (int, float)) and qty > 0 else Decimal("1")
        except InvalidOperation:
            cantidad_dec = Decimal("1")

        try:
            v_unit_dec = Decimal(str(v_unit_val)) if v_unit_val is not None else Decimal("0")
        except InvalidOperation:
            v_unit_dec = Decimal("0")

        if v_unit_dec <= 0 and precio_total > 0 and cantidad_dec > 0:
            v_unit_dec = precio_total / cantidad_dec

        # Si no tiene costo unitario ni precio total, es un título/texto de encabezado -> IGNORAR
        if v_unit_dec <= 0 and precio_total <= 0:
            continue

        unidad_str = str(celdas[2]).strip() if len(celdas) > 2 and celdas[2] is not None else "Global"

        seccion_final = "MANO_DE_OBRA" if tipo == "mano_obra" else seccion_actual

        receta_filas.append({
            "producto": nombre_producto,
            "nombre_material": col_a_str[:150],
            "cantidad_base": cantidad_dec,
            "v_unit": v_unit_dec,
            "unidad": unidad_str,
            "seccion": seccion_final,
            "es_mano_obra": (tipo == "mano_obra")
        })

    return {
        "filas": receta_filas,
        "precio_costo": precio_costo,
        "precio_venta": precio_venta,
        "precio_con_iva": precio_con_iva,
    }


def get_or_create_unidad(db, nombre_excel: str) -> UnidadMedida:
    norm = nombre_excel.upper().strip()
    abrev, nombre_completo = UNIDADES_NORMALIZADAS.get(norm, (norm.lower()[:10], norm[:50]))

    u = db.query(UnidadMedida).filter(
        (UnidadMedida.nombre.ilike(nombre_completo)) |
        (UnidadMedida.abreviatura == abrev)
    ).first()
    if not u:
        u = UnidadMedida(nombre=nombre_completo, abreviatura=abrev)
        db.add(u)
        db.flush()
    return u


def get_or_create_material(db, nombre: str, v_unit: Decimal, um_id: int) -> Material:
    m = db.query(Material).filter(Material.nombre == nombre).first()
    if not m:
        m = Material(nombre=nombre, costo_base=v_unit, unidad_medida_id=um_id, activo=True)
        db.add(m)
        db.flush()
    else:
        if v_unit > 0:
            m.costo_base = v_unit
            db.flush()
    return m


def main():
    parser = argparse.ArgumentParser(description="Importador de estructuras de costos desde Excel")
    parser.add_argument("--dry-run", action="store_true", help="Solo analiza y muestra estadísticas sin guardar en la BD")
    args = parser.parse_args()

    print(f"📂 Leyendo Excel: {EXCEL_PATH}")
    if not EXCEL_PATH.exists():
        print(f"❌ No se encontró el archivo: {EXCEL_PATH}")
        sys.exit(1)

    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    print(f"Total hojas a procesar: {len(wb.sheetnames)}")

    if args.dry_run:
        print("🔍 MODO DRY-RUN ACTIVADO (No se modificará la base de datos)\n")
        total_filas = 0
        conteo_secciones = {}
        hojas_procesadas = 0

        for sheet_name in wb.sheetnames:
            nombre = limpiar_nombre_hoja(sheet_name)
            res = parsear_hoja(wb[sheet_name], nombre)
            filas = res["filas"]
            if not filas:
                continue

            hojas_procesadas += 1
            total_filas += len(filas)
            for f in filas:
                sec = f["seccion"]
                conteo_secciones[sec] = conteo_secciones.get(sec, 0) + 1

            if hojas_procesadas in (1, 2, 3, len(wb.sheetnames)):
                print(f"  • Hoja [{sheet_name.strip()}]: {len(filas)} ítems identificados")
                for f in filas[:4]:
                    print(f"      [{f['seccion']:<12}] {f['nombre_material']:<40} Qty: {f['cantidad_base']} ${f['v_unit']}")

        print(f"\n📊 RESUMEN DRY-RUN:")
        print(f"  Hojas procesadas: {hojas_procesadas} / {len(wb.sheetnames)}")
        print(f"  Total ítems de receta: {total_filas}")
        print("  Distribución por sección:")
        for sec, cnt in sorted(conteo_secciones.items()):
            print(f"    - {sec:<15}: {cnt} filas")
        return

    # Modo REAL (escritura en base de datos)
    db = session_local()
    try:
        tipo_cama = db.query(TipoProducto).filter(TipoProducto.nombre.ilike("%cama%")).first()
        if not tipo_cama:
            tipo_cama = TipoProducto(nombre="Cama", descripcion="Camas y muebles de dormitorio")
            db.add(tipo_cama)
            db.flush()

        # Limpiar importaciones previas asociadas al Excel
        prev = db.query(Producto).filter(Producto.hoja_excel.isnot(None)).all()
        for p in prev:
            db.query(ProductoMaterial).filter(ProductoMaterial.producto_id == p.id).delete()
            db.delete(p)
        db.flush()
        print(f"🧹 Eliminados {len(prev)} productos anteriores del Excel\n")

        creados = 0
        total_recetas_creadas = 0

        for sheet_name in wb.sheetnames:
            nombre_prod = limpiar_nombre_hoja(sheet_name)
            datos = parsear_hoja(wb[sheet_name], nombre_prod)
            filas = datos["filas"]
            if not filas and not datos["precio_costo"]:
                continue

            pv = datos["precio_venta"]
            piva = datos["precio_con_iva"]
            costo = datos["precio_costo"]

            if pv and not piva:
                piva = (pv * Decimal("1.16")).quantize(Decimal("0.01"))
            if piva and not pv:
                pv = (piva / Decimal("1.16")).quantize(Decimal("0.01"))
            if not costo and pv:
                costo = (pv / Decimal("1.30")).quantize(Decimal("0.01"))

            producto = Producto(
                nombre=nombre_prod,
                tipo_producto_id=tipo_cama.id,
                activo=True,
                hoja_excel=sheet_name,
                ancho_base=Decimal("1.60"),
                largo_base=Decimal("1.90"),
                precio_costo_base=costo,
                precio_venta_base=pv,
                precio_venta_con_iva=piva,
            )
            db.add(producto)
            db.flush()

            insumos_insertados = 0
            for f in filas:
                um = get_or_create_unidad(db, f["unidad"])
                mat = get_or_create_material(db, f["nombre_material"], f["v_unit"], um.id)

                # Fix Bug 2: verificar duplicados por (producto_id, material_id, seccion)
                existente = db.query(ProductoMaterial).filter(
                    ProductoMaterial.producto_id == producto.id,
                    ProductoMaterial.material_id == mat.id,
                    ProductoMaterial.seccion == f["seccion"]
                ).first()

                if existente:
                    # Si ya existe en la misma sección, se acumula la cantidad
                    existente.cantidad_base += f["cantidad_base"]
                else:
                    pm = ProductoMaterial(
                        producto_id=producto.id,
                        material_id=mat.id,
                        seccion=f["seccion"],
                        cantidad_base=f["cantidad_base"],
                        tipo_escala="FIJO",
                        observaciones=f"hoja:{sheet_name}"
                    )
                    db.add(pm)
                    insumos_insertados += 1

            total_recetas_creadas += insumos_insertados
            creados += 1
            print(f"  ✅ [{sheet_name.strip()[:30]:<30}] -> Prod #{producto.id} | {insumos_insertados} recetas | Costo: {costo}")

        db.commit()
        print(f"\n🎉 ¡Importación completada exitosamente!")
        print(f"   - Productos creados: {creados}")
        print(f"   - Recetas (ProductoMaterial) insertadas: {total_recetas_creadas}")

    except Exception as e:
        db.rollback()
        print(f"\n❌ Error en la importación: {e}")
        traceback.print_exc()
        sys.exit(1)
    finally:
        db.close()


if __name__ == "__main__":
    main()
